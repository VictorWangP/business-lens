"""Import user-supplied tabular values into an immutable, typed DuckDB dataset."""
import csv
import io
import re
import shutil
import uuid
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from app.config import settings

MAX_ROWS = 100_000
MAX_COLUMNS = 200
NUMBER = re.compile(r'^[+-]?\d+(?:\.\d+)?$')
ISO_DATE = re.compile(r'^\d{4}-\d{2}-\d{2}$')


def _rows(filename: str, content: bytes, sheet: str | None):
    extension = Path(filename).suffix.lower()
    if extension == '.csv':
        try:
            text = content.decode('utf-8-sig')
        except UnicodeDecodeError as exc:
            raise ValueError('CSV 必须使用 UTF-8 编码，请重新导出文件') from exc
        try:
            dialect = csv.Sniffer().sniff(text[:16384], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        return csv.reader(io.StringIO(text, newline=''), dialect, strict=True), None, [], None
    if extension == '.xlsx':
        from openpyxl import load_workbook
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(f.file_size for f in archive.infolist()) > 200 * 1024 * 1024:
                    raise ValueError('Excel 解压后超过 200MB，请拆分文件')
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        except (zipfile.BadZipFile, KeyError, OSError) as exc:
            raise ValueError('Excel 文件无效，请使用标准 .xlsx 文件') from exc
        selected = sheet or workbook.sheetnames[0]
        if selected not in workbook.sheetnames:
            workbook.close()
            raise ValueError('指定的工作表不存在')
        worksheets = workbook.sheetnames
        def values():
            for row in workbook[selected].iter_rows():
                if len(row) > MAX_COLUMNS:
                    raise ValueError(f'列数不能超过 {MAX_COLUMNS}')
                if any(cell.data_type in {'f', 'e'} for cell in row):
                    raise ValueError('Excel 包含公式或错误单元格，请先复制并粘贴为值后导入')
                yield [cell.value for cell in row]
        return values(), selected, worksheets, workbook.close
    raise ValueError('业务数据导入支持 .csv 和 .xlsx')


def _text(value):
    return value.isoformat() if isinstance(value, (date, datetime)) else str(value)


def _column(values):
    present = [v for v in values if v is not None]
    if not present:
        return 'VARCHAR', values
    if all(isinstance(v, bool) for v in present):
        return 'BOOLEAN', values
    if all(isinstance(v, (date, datetime)) or (isinstance(v, str) and ISO_DATE.fullmatch(v)) for v in present):
        try:
            if any(isinstance(v, datetime) for v in present):
                return 'TIMESTAMP', [None if v is None else (datetime.combine(v, datetime.min.time())
                    if isinstance(v, date) and not isinstance(v, datetime) else datetime.fromisoformat(v)
                    if isinstance(v, str) else v) for v in values]
            return 'DATE', [None if v is None else date.fromisoformat(v) if isinstance(v,str) else v for v in values]
        except ValueError:
            pass
    # Preserve identifier-like values such as 0012 as text.
    texts = [_text(v) for v in present]
    numeric = all(not isinstance(v, bool) and NUMBER.fullmatch(t) and not re.match(r'^[+-]?0\d', t)
                  for v,t in zip(present,texts))
    if numeric:
        try:
            decimals = [None if v is None else Decimal(str(v)) for v in values]
            valid = [v for v in decimals if v is not None]
            scale = max(max(-v.as_tuple().exponent, 0) for v in valid)
            integer_digits = max(max(v.adjusted()+1,1) for v in valid)
            if integer_digits + scale <= 38:
                return f'DECIMAL(38,{scale})', decimals
        except (InvalidOperation, ValueError):
            pass
    return 'VARCHAR', [None if v is None else _text(v) for v in values]


def import_file(user_id: str, filename: str, content: bytes, sheet: str | None = None) -> dict:
    if not content or len(content) > settings.MAX_UPLOAD_BYTES:
        raise ValueError('文件为空或超过上传大小限制')
    rows, selected_sheet, sheets, close = _rows(filename, content, sheet)
    try:
        header = next(rows, None)
        if not header or len(header) > MAX_COLUMNS:
            raise ValueError(f'文件必须包含表头，列数不超过 {MAX_COLUMNS}')
        names = [str(v).strip() if v is not None else '' for v in header]
        if any(not name or len(name)>128 or '\x00' in name for name in names):
            raise ValueError('表头不能为空、含空字符或超过 128 字符')
        if len({n.casefold() for n in names}) != len(names):
            raise ValueError('表头重复，请使用不同列名')
        records = []
        for line, row in enumerate(rows, 2):
            if not row or all(v is None or v == '' for v in row):
                continue
            if len(row) != len(names):
                raise ValueError(f'第 {line} 行列数与表头不一致')
            if len(records) >= MAX_ROWS:
                raise ValueError(f'单文件最多 {MAX_ROWS} 行，请拆分文件')
            records.append([None if v is None or v == '' else v for v in row])
        if not records:
            raise ValueError('文件没有业务数据行')
        types, converted = zip(*[_column([r[i] for r in records]) for i in range(len(names))])
    except csv.Error as exc:
        raise ValueError('CSV 格式无效，请检查引号及分隔符') from exc
    finally:
        if close:
            close()
    source_id = str(uuid.uuid4())
    directory = settings.USER_STORAGE_DIR / user_id / 'data_sources' / source_id
    directory.mkdir(parents=True, exist_ok=False)
    path = directory / 'business.duckdb'
    try:
        import duckdb
        with duckdb.connect(str(path), config={'threads':'2','memory_limit':'256MB'}) as conn:
            definitions = ','.join('"'+n.replace('"','""')+'" '+t for n,t in zip(names,types))
            conn.execute(f'CREATE TABLE _records ({definitions})')
            # Arrow ingestion keeps Decimal/date types, avoiding per-row SQL conversion and round trips.
            import pyarrow as pa
            arrow = pa.table({name: list(values) for name, values in zip(names,converted)})
            conn.register('import_values', arrow)
            conn.execute('INSERT INTO _records SELECT * FROM import_values')
            conn.unregister('import_values')
            conn.execute('CREATE VIEW dataset AS SELECT * FROM _records')
        original = directory / ('original' + Path(filename).suffix.lower())
        original.write_bytes(content)
    except BaseException:
        shutil.rmtree(directory)
        raise
    return {'id': source_id, 'path': path, 'row_count': len(records), 'sheet': selected_sheet, 'sheets': sheets,
            'columns': [{'table_name':'dataset','column_name':n,'data_type':t} for n,t in zip(names,types)]}
