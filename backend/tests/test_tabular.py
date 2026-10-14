import json
"""Real CSV/XLSX -> persisted DuckDB -> authorized SQL, without a database server or LLM."""
import io
from datetime import date
from decimal import Decimal
from types import SimpleNamespace
import httpx
import pytest
from openpyxl import Workbook
from app.application.tabular import import_file
from app.connectors.registry import execute, inspect_source
from app.infrastructure.credentials import encrypt
from app.analytics.presentation import json_value
from app.database import AsyncSessionLocal
from app.models.platform import DataSource
from app.agents.executor import AgentExecutor
from main import app
from test_platform import seed, auth

CSV = 'order_date,region,amount,customer_id\n2026-01-01,华东,0.10,0012\n2026-01-02,华东,0.20,0013\n2026-02-01,华北,100.00,0014\n'.encode()

def source(result, user='local-user'):
    return SimpleNamespace(kind='duckdb', user_id=user, connection_secret=encrypt('duckdb://' + str(result['path'])),
                           schema_name='main', allowed_views=['dataset'])

@pytest.mark.asyncio
async def test_real_csv_decimal_date_queries_and_filesystem_isolation():
    result = import_file('local-user', 'orders.csv', CSV)
    s = source(result)
    assert result['row_count'] == 3
    assert len(await inspect_source(s)) == 4
    values = await execute(s, 'SELECT SUM(amount) AS total FROM dataset WHERE order_date >= :start_date AND order_date < :end_date',
                           {'start_date':date(2026,1,1), 'end_date':date(2026,2,1)})
    assert values['rows'] == [{'total':Decimal('0.30')}]
    assert json_value(values)['rows'][0]['total'] == '0.30'
    assert (await execute(s, 'SELECT customer_id FROM dataset ORDER BY customer_id'))['rows'][0]['customer_id'] == '0012'
    for sql in ["SELECT * FROM read_csv('/etc/passwd')", 'SELECT * FROM _records', 'DELETE FROM dataset',
                "SELECT current_setting('enable_external_access') FROM dataset", 'SELECT * FROM information_schema.tables']:
        with pytest.raises(ValueError): await execute(s, sql)
    s.user_id = 'other-user'
    with pytest.raises(ValueError): await execute(s, 'SELECT * FROM dataset')

@pytest.mark.asyncio
async def test_upload_api_metric_and_agent_use_actual_imported_data():
    a,b,_ = await seed()
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url='http://test') as client:
        r = await client.post('/api/data-sources/import', headers=auth(a), data={'name':'真实订单'}, files={'file':('orders.csv', CSV, 'text/csv')})
        assert r.status_code == 200, r.text
        s = r.json(); sid=s['id']
        assert s['kind']=='duckdb' and s['row_count']==3
        assert 'connection_secret' not in s and 'path' not in s
        assert (await client.post(f'/api/data-sources/{sid}/query',headers=auth(b),json={'sql':'SELECT * FROM dataset'})).status_code==404
        r = await client.post(f'/api/data-sources/{sid}/metrics', headers=auth(a), json={'name':'销售额','description':'按订单日期筛选，未含税调整，原始币种',
            'sql_template':'SELECT SUM(amount) AS total FROM dataset WHERE order_date>=:start_date AND order_date<:end_date','unit':'元'})
        assert r.status_code==200, r.text
        mid=r.json()['id']
        agent = AgentExecutor(a, [], [sid], {})
        tool = next(t for t in agent.tools() if t.name=='query_metrics')
        output = await tool.ainvoke({'metric_id':mid,'start_date':'2026-01-01','end_date':'2026-02-01'})
        assert 'error' not in json.loads(output)
        assert agent.queries[0]['rows']==[{'total':'0.30'}]
        assert agent.queries[0]['citation_id'].startswith('Q')
        assert (await client.delete(f'/api/data-sources/{sid}',headers=auth(a))).status_code==204
        with pytest.raises(ValueError): await agent.source(sid)


def test_xlsx_sheet_and_formula_validation():
    book=Workbook(); book.active.title='订单'
    book.active.append(['日期','金额']); book.active.append([date(2026,1,1),1.25])
    other=book.create_sheet('客户'); other.append(['客户']); other.append(['客户甲'])
    data=io.BytesIO(); book.save(data)
    result=import_file('xlsx-user','orders.xlsx',data.getvalue(),'订单')
    assert result['row_count']==1 and result['sheet']=='订单' and result['sheets']==['订单','客户']
    book.active.append([date(2026,1,2),'=1+2']); data=io.BytesIO(); book.save(data)
    with pytest.raises(ValueError,match='公式'): import_file('xlsx-user','formulas.xlsx',data.getvalue())
    with pytest.raises(ValueError,match='工作表不存在'): import_file('xlsx-user','orders.xlsx',data.getvalue(),'missing')

@pytest.mark.parametrize('content', [b'a,a\n1,2', b'a,b\n1,2,3', b'a,b\n', b',b\n1,2'])
def test_malformed_csv_not_imported(content):
    with pytest.raises(ValueError): import_file('invalid-user','invalid.csv',content)
